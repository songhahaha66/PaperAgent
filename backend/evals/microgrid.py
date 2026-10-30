"""Offline checks for the supplied microgrid question 1, never a runtime gate.

Assumptions: each source row represents ten minutes, both charge and discharge
have efficiency 0.9, curtailment is allowed, export is not. Alternative conventions
need a separate reference. Matching the objective alone does not prove feasibility.
"""
import json
from pathlib import Path
import numpy as np
import openpyxl
from scipy.optimize import linprog
from scipy.sparse import lil_matrix


def grade_result1(source: Path, result: Path):
    book=openpyxl.load_workbook(source,read_only=True,data_only=True)
    try:data=np.array([row[1:4] for row in list(book.active.values)[1:]],float)
    finally:book.close()
    price,load,pv=data.T;n=len(price);dt=1/6;eta=.9
    if n!=144 or not np.isfinite(data).all():raise ValueError('需要完整的144时段原始数据')
    # Independent power variables: grid, charge, discharge, curtailment, E0..En.
    A=lil_matrix((2*n,5*n+1));rhs=np.r_[load-pv,np.zeros(n)]
    for t in range(n):
        A[t,[t,n+t,2*n+t,3*n+t]]=[1,-1,1,-1]
        A[n+t,[n+t,2*n+t,4*n+t,4*n+t+1]]=[-eta*dt,dt/eta,-1,1]
    bounds=[(0,None)]*n+[(0,5000)]*(2*n)+[(0,float(v))for v in pv]+[(1200,10800)]*(n+1)
    bounds[4*n]=(6000,6000);bounds[-1]=(6000,6000)
    optimum=linprog(np.r_[price*dt,np.zeros(4*n+1)],A_eq=A.tocsr(),b_eq=rhs,bounds=bounds,method='highs')
    if not optimum.success:raise ValueError(optimum.message)
    checks={};details={'reference_cost_yuan':float(optimum.fun)}
    book=openpyxl.load_workbook(result,read_only=True,data_only=True)
    try:
        if not {'计划购电量','充放电量'}<=set(book.sheetnames):
            return {'passed':False,'checks':{'required_sheets':False},'details':details}
        grid=[r[1]for r in list(book['计划购电量'].values)[1:]]
        checks['144_finite_nonnegative_intervals']=len(grid)==144 and all(
            isinstance(v,(int,float))and np.isfinite(v)and v>=-1e-6 for v in grid)
        if checks['144_finite_nonnegative_intervals']:
            cost=float(np.dot(grid,price));details['actual_cost_yuan']=cost
            # Permit ordinary rounding; never require one particular optimal schedule.
            checks['objective_within_2_yuan']=abs(cost-optimum.fun)<=2
        rows=list(book['充放电量'].values)
        try:
            ch=np.array([r[1]for r in rows[1:7]],float)
            dis=np.array([r[2]for r in rows[1:7]],float)
            start,end=float(rows[1][4]),float(rows[2][4])
            checks['terminal_energy']=abs(start-6000)<.1 and abs(end-6000)<.1
            checks['four_hour_power_bound']=bool(np.isfinite(ch).all()and np.isfinite(dis).all()
                and np.min(ch)>=0 and np.min(dis)>=0 and max(max(ch),max(dis))<=20000.1)
            residual=float(end-start-eta*sum(ch)+sum(dis)/eta)
            details['daily_energy_residual_kwh']=residual
            checks['independent_daily_energy_balance']=np.isfinite(residual)and abs(residual)<.2
        except (TypeError,ValueError,IndexError):checks['storage_template_values']=False
    finally:book.close()
    return {'passed':bool(checks)and all(bool(v)for v in checks.values()),
            'checks':{k:bool(v)for k,v in checks.items()},'details':details,
            'scope':'问题1目标值、模板数值与日能量守恒；不代替逐时段可行性或问题2—4验证。'}


if __name__=='__main__':
    import sys
    result=grade_result1(Path(sys.argv[1]),Path(sys.argv[2]))
    print(json.dumps(result,ensure_ascii=False,indent=2))
    sys.exit(0 if result['passed'] else 1)
