"""Single docx rendering engine: pristine template copy + PaperIR → document.

Every IR block maps to a deterministic OOXML operation so that rendering the same
IR against the same template yields the same file (see `normalize_docx_bytes`).
"""

from __future__ import annotations

import shutil
import zipfile
from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph as DocxParagraph

from ..schemas.paper_ir import (
    Citation,
    Code,
    Equation,
    FigureRef,
    ListBlock,
    PaperIR,
    Paragraph,
    TableRows,
)
from ..schemas.template_spec import Slot, TemplateSpec

VOLATILE_PARTS = {"docProps/core.xml", "docProps/app.xml"}
RSID_LOCAL_NAMES = {"rsidR", "rsidRDefault", "rsidP", "rsidRPr", "rsidTr", "rsidSect", "docId"}
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WP_INLINE = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
REMOVE_ROLES = {"instruction_delete", "example_delete"}
REPLACE_ROLES = {"placeholder_fill", "caption"}
APPEND_ROLES = {"figure_slot"}
CODE_FONT = "Consolas"
try:
    if "m" not in nsmap:
        nsmap["m"] = MATH_NS
except Exception:
    pass


def render_docx(
    spec: TemplateSpec,
    ir: PaperIR,
    template_path: Path | str | None,
    output_path: Path | str,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    template = Path(template_path) if template_path else None
    if template and template.exists():
        if template.resolve() != output_path.resolve():
            shutil.copyfile(template, output_path)
        document = Document(str(output_path))
        _fill_from_template(document, spec, ir, template.parent)
    else:
        document = Document()
        _fill_blank_document(document, spec, ir, output_path.parent)
    _append_references(document, ir)
    document.save(str(output_path))
    return output_path


def render_docx_from_ir(spec: TemplateSpec, ir: PaperIR, output_path: Path | str) -> Path:
    return render_docx(spec, ir, None, output_path)


def normalize_docx_bytes(path: Path | str) -> list[tuple[str, bytes]]:
    """Byte snapshot excluding volatile core properties and revision IDs."""
    with zipfile.ZipFile(path) as archive:
        names = sorted(name for name in archive.namelist() if name not in VOLATILE_PARTS)
        return [(name, _strip_volatile_xml(name, archive.read(name))) for name in names]


# ---------------------------------------------------------------------------
# Template filling
# ---------------------------------------------------------------------------


def index_body_blocks(document: Document) -> dict[str, object]:
    """Map parser block ids (P###/G###/T###) to body elements, mirroring `parse_docx`."""
    mapping: dict[str, object] = {}
    paragraph_index = table_index = figure_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            mapping[f"T{table_index:03d}"] = child
            table_index += 1
            continue
        if child.tag != qn("w:p"):
            continue
        text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
        has_image = any(True for _ in child.iter(WP_INLINE))
        if has_image and not text:
            mapping[f"G{figure_index:03d}"] = child
            figure_index += 1
        else:
            mapping[f"P{paragraph_index:03d}"] = child
            paragraph_index += 1
    return mapping


def _fill_from_template(document: Document, spec: TemplateSpec, ir: PaperIR, workspace: Path) -> None:
    block_text = {block.id: block.text.strip() for block in spec.blocks}
    by_id = index_body_blocks(document)
    claimed: set[int] = set()

    def resolve(slot: Slot):
        element = by_id.get(slot.anchor_block)
        expected = block_text.get(slot.anchor_block)
        if element is not None and element.tag == qn("w:p") and expected:
            actual = "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()
            if actual != expected:
                element = None
        if element is None and expected:
            # Spec built from a different revision of the template: fall back to text.
            for paragraph in document.paragraphs:
                if id(paragraph._element) in claimed or paragraph.text.strip() != expected:
                    continue
                element = paragraph._element
                break
        if element is not None:
            claimed.add(id(element))
        return element

    removals = []
    fills: list[tuple[object, Slot]] = []
    for slot in spec.slots:
        if slot.role in REMOVE_ROLES:
            element = resolve(slot)
            if element is not None and element.tag == qn("w:p"):
                removals.append(element)
        elif slot.role in REPLACE_ROLES | APPEND_ROLES:
            section = ir.sections.get(slot.id)
            if not section or not section.blocks:
                continue
            element = resolve(slot)
            if element is not None and element.tag == qn("w:p"):
                fills.append((element, slot))

    _fill_tables(document, spec, ir, by_id)

    for element, slot in fills:
        anchor = DocxParagraph(element, document._body)
        blocks = ir.sections[slot.id].blocks
        if slot.role in REPLACE_ROLES:
            _clear_paragraph(anchor)
            cursor = _write_block(anchor, blocks[0], ir, workspace, primary=True)
            rest = blocks[1:]
        else:
            cursor = anchor
            rest = blocks
        for block in rest:
            cursor = _write_block(_insert_after(cursor, template=anchor), block, ir, workspace, primary=True)

    for element in removals:
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def _fill_blank_document(document: Document, spec: TemplateSpec, ir: PaperIR, workspace: Path) -> None:
    for slot in spec.slots:
        if slot.role in REMOVE_ROLES:
            continue
        if slot.role == "heading":
            title = slot.title or (slot.section_path[-1] if slot.section_path else slot.id)
            _add_heading(document, title, level=min(len(slot.section_path) or 1, 3))
            continue
        if slot.role == "fixed_text" and slot.title:
            document.add_paragraph(slot.title)
            continue
        section = ir.sections.get(slot.id)
        if not section:
            continue
        for block in section.blocks:
            _write_block(document.add_paragraph(), block, ir, workspace, primary=True)


def _append_references(document: Document, ir: PaperIR) -> None:
    if not ir.references:
        return
    existing = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "参考文献" in existing:
        return
    _add_heading(document, "参考文献", level=1)
    for index, reference in enumerate(ir.references.values(), start=1):
        document.add_paragraph(f"[{index}] {reference.text}")


def _add_heading(document: Document, text: str, level: int) -> None:
    try:
        document.add_heading(text, level=level)
    except KeyError:
        # Templates without built-in "Heading N" styles: fall back to a bold paragraph.
        paragraph = document.add_paragraph()
        paragraph.add_run(text).bold = True


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------


def _fill_tables(document: Document, spec: TemplateSpec, ir: PaperIR, by_id: dict[str, object]) -> None:
    tables_by_element = {id(table._tbl): table for table in document.tables}
    fallback = iter(document.tables)
    for slot in (slot for slot in spec.slots if slot.role == "table"):
        section = ir.sections.get(slot.id)
        if not section:
            continue
        table_block = next((block for block in section.blocks if isinstance(block, TableRows)), None)
        if table_block is None:
            continue
        element = by_id.get(slot.anchor_block)
        table = tables_by_element.get(id(element)) if element is not None else None
        if table is None:
            table = next(fallback, None)
        if table is None:
            continue
        _write_table(table, table_block)


def _write_table(table: Table, table_block: TableRows) -> None:
    rows = table_block.rows
    if not rows or not table.rows:
        return
    while len(table.rows) < len(rows):
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    # Drop leftover template rows (typically example rows) but always keep the header.
    for surplus in list(table.rows)[max(len(rows), 1):]:
        table._tbl.remove(surplus._tr)
    for row_index, values in enumerate(rows):
        cells = table.rows[row_index].cells
        for col_index, value in enumerate(values[: len(cells)]):
            _set_cell_text(cells[col_index], "" if value is None else str(value))


def _set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    r_pr = next((deepcopy(run._element.rPr) for run in paragraph.runs if run._element.rPr is not None), None)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    _clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if r_pr is not None:
        run._element.insert(0, r_pr)


# ---------------------------------------------------------------------------
# Blocks
# ---------------------------------------------------------------------------


def _write_block(paragraph: DocxParagraph, block, ir: PaperIR, workspace: Path, primary: bool) -> DocxParagraph:
    """Write one IR block starting at `paragraph`; return the last paragraph written."""
    if isinstance(block, Paragraph):
        _set_text(paragraph, block.text)
        return paragraph
    if isinstance(block, Code):
        run = _set_text(paragraph, block.text)
        run.font.name = CODE_FONT
        run.font.size = Pt(10)
        try:
            run._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
        except Exception:
            pass
        return paragraph
    if isinstance(block, ListBlock):
        cursor = paragraph
        for index, item in enumerate(block.items):
            if index:
                cursor = _insert_after(cursor, template=paragraph)
            prefix = f"{index + 1}. " if block.ordered else "• "
            _set_text(cursor, prefix + item)
        return cursor
    if isinstance(block, FigureRef):
        return _write_figure(paragraph, block, ir, workspace)
    if isinstance(block, Equation):
        _write_equation(paragraph, block.latex)
        return paragraph
    if isinstance(block, Citation):
        reference = ir.references.get(block.ref_id)
        _set_text(paragraph, reference.text if reference else f"[{block.ref_id}]")
        return paragraph
    if isinstance(block, TableRows):
        _set_text(paragraph, "\n".join("\t".join(row) for row in block.rows))
        return paragraph
    return paragraph


def _write_figure(paragraph: DocxParagraph, block: FigureRef, ir: PaperIR, workspace: Path) -> DocxParagraph:
    artifact = ir.artifacts.get(block.artifact_id)
    path = Path(artifact.path) if artifact else None
    if path and not path.is_absolute():
        candidates = [workspace / path, workspace.parent / path]
        path = next((candidate for candidate in candidates if candidate.exists()), candidates[0])
    if path and path.exists():
        _clear_paragraph(paragraph)
        run = paragraph.add_run()
        try:
            run.add_picture(str(path), width=Inches(4.8))
        except Exception:
            _set_text(paragraph, block.caption or block.artifact_id)
            return paragraph
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if block.caption:
            caption = _insert_after(paragraph, template=paragraph)
            _set_text(caption, block.caption)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return caption
        return paragraph
    _set_text(paragraph, block.caption or f"[图:{block.artifact_id}]")
    return paragraph


def _write_equation(paragraph: DocxParagraph, latex: str) -> None:
    omath = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = latex
    run.append(text)
    omath.append(run)
    paragraph._element.append(omath)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ---------------------------------------------------------------------------
# Paragraph helpers
# ---------------------------------------------------------------------------


def _clear_paragraph(paragraph: DocxParagraph) -> None:
    """Remove runs/hyperlinks but keep the paragraph properties (style, spacing)."""
    for child in list(paragraph._element):
        if child.tag != qn("w:pPr"):
            paragraph._element.remove(child)


def _set_text(paragraph: DocxParagraph, text: str):
    """Replace paragraph text with a single run that inherits the first run's rPr."""
    r_pr = next((deepcopy(run._element.rPr) for run in paragraph.runs if run._element.rPr is not None), None)
    _clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if r_pr is not None:
        run._element.insert(0, r_pr)
    return run


def _insert_after(paragraph: DocxParagraph, template: DocxParagraph | None = None) -> DocxParagraph:
    """Insert an empty paragraph after `paragraph`, cloning `template`'s pPr and first rPr."""
    template = template or paragraph
    new_p = paragraph._element.makeelement(qn("w:p"), {})
    p_pr = template._element.find(qn("w:pPr"))
    if p_pr is not None:
        new_p.append(deepcopy(p_pr))
    paragraph._element.addnext(new_p)
    added = DocxParagraph(new_p, paragraph._parent)
    r_pr = next((deepcopy(run._element.rPr) for run in template.runs if run._element.rPr is not None), None)
    if r_pr is not None:
        added.add_run("")._element.insert(0, r_pr)
    return added


def _strip_volatile_xml(name: str, data: bytes) -> bytes:
    if not name.endswith(".xml"):
        return data
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return data
    for node in list(root.iter()):
        for attr in list(node.attrib):
            local = attr.rsplit("}", 1)[-1]
            if local in RSID_LOCAL_NAMES or local.lower().startswith("rsid"):
                del node.attrib[attr]
    _remove_rsids(root)
    return ET.tostring(root, encoding="utf-8")


def _remove_rsids(root: ET.Element) -> None:
    for parent in root.iter():
        for child in list(parent):
            if child.tag.rsplit("}", 1)[-1] == "rsids":
                parent.remove(child)
