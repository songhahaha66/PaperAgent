from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from ..schemas.paper_ir import PaperIR, Paragraph, SectionContent, TableRows
from ..schemas.template_spec import TemplateSpec
from ..template.spec_builder import build_template_spec
from .harness import EvalCase


FILL_TEXT = "本节给出研究背景、操作步骤与预期结果，内容覆盖模板占位要求并超过最小字数。"


def build_golden_cases(workdir: Path) -> list[EvalCase]:
    workdir = Path(workdir)
    workdir.mkdir(parents=True, exist_ok=True)
    cases: list[EvalCase] = []
    for name, builder in (
        ("course-a", _course_a),
        ("course-b", _course_b),
        ("thesis-a", _thesis_a),
        ("thesis-b", _thesis_b),
        ("journal-a", _journal_a),
        ("journal-b", _journal_b),
    ):
        template = builder(workdir / f"{name}.docx")
        spec = build_template_spec(template, template_id=abs(hash(name)) % 10_000)
        ir = _filled_ir(spec, name)
        cases.append(
            EvalCase(
                name=name,
                spec=spec,
                ir=ir,
                template_path=template,
                output_mode="word",
            )
        )
    return cases


def _filled_ir(spec: TemplateSpec, work_id: str) -> PaperIR:
    sections: dict[str, SectionContent] = {}
    for slot in spec.slots:
        if slot.role == "placeholder_fill":
            sections[slot.id] = SectionContent(
                slot_id=slot.id,
                blocks=[Paragraph(text=f"{slot.title or slot.id}。{FILL_TEXT}")],
            )
        elif slot.role == "table":
            sections[slot.id] = SectionContent(
                slot_id=slot.id,
                blocks=[TableRows(table_id=slot.anchor_block, rows=[["项目", "结果"], ["吞吐", "120"]])],
            )
        elif slot.role == "caption":
            sections[slot.id] = SectionContent(
                slot_id=slot.id,
                blocks=[Paragraph(text=slot.title or "图1 实验结果")],
            )
    return PaperIR(work_id=work_id, revision=1, sections=sections)


def _add_outline_style(document: Document, name: str, level: int, size_pt: int = 16) -> None:
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    style.font.size = Pt(size_pt)
    p_pr = style.element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def _course_a(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "一级标题", 0, 18)
    _add_outline_style(doc, "二级标题", 1, 14)
    doc.add_paragraph("1 实验目的", style="一级标题")
    doc.add_paragraph("1.1 背景说明", style="二级标题")
    doc.add_paragraph("学号：20260001")
    doc.add_paragraph("请在此处填写实验步骤")
    doc.add_paragraph("例如：print(\"hello\")")
    doc.add_paragraph("成稿后删除：格式要求使用宋体小四")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "示例行"
    table.cell(1, 1).text = "—"
    doc.save(str(path))
    return path


def _course_b(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "一级标题", 0, 18)
    doc.add_paragraph("1 实验环境", style="一级标题")
    doc.add_paragraph("2 实验过程", style="一级标题")
    doc.add_paragraph("请在此处填写实验过程")
    doc.add_paragraph("例如：SELECT * FROM t")
    doc.add_paragraph("成稿后删除：写作说明可删")
    doc.save(str(path))
    return path


def _thesis_a(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "章标题", 0, 18)
    for title in ("摘要", "绪论", "方法", "实验", "结论"):
        doc.add_paragraph(title, style="章标题")
        doc.add_paragraph(f"请在此处填写{title}")
    doc.add_paragraph("例如：样例段落不要保留")
    doc.save(str(path))
    return path


def _thesis_b(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "章标题", 0, 18)
    for title in ("中文摘要", "研究背景", "系统设计", "实验验证", "总结与展望"):
        doc.add_paragraph(title, style="章标题")
        doc.add_paragraph(f"请填写{title}正文")
    doc.add_paragraph("成稿后删除：格式要求使用宋体小四")
    doc.save(str(path))
    return path


def _journal_a(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "HeadingCN", 0, 16)
    for title in ("Abstract", "Introduction", "Methods", "Results"):
        doc.add_paragraph(title, style="HeadingCN")
        doc.add_paragraph(f"请在此处填写{title}")
    doc.add_paragraph("例如：dummy result")
    doc.save(str(path))
    return path


def _journal_b(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "HeadingCN", 0, 16)
    for title in ("摘要", "引言", "方法", "结果与讨论"):
        doc.add_paragraph(title, style="HeadingCN")
        doc.add_paragraph(f"请在此处填写{title}")
    doc.add_paragraph("图1 实验结果")
    doc.add_paragraph("例如：placeholder citation")
    doc.save(str(path))
    return path
