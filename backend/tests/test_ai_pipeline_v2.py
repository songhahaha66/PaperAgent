import asyncio
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from ai_system.core_managers.stream_manager import StreamOutputManager
from ai_system.judge.heuristic import HeuristicJudge, infer_role
from ai_system.render.docx_renderer import normalize_docx_bytes, render_docx
from ai_system.render.markdown_renderer import render_markdown
from ai_system.runtime.events import EventEmitter
from ai_system.sandbox.runner import run_sandbox
from ai_system.schemas.paper_ir import PaperIR, Paragraph, SectionContent, TableRows
from ai_system.schemas.plan import derive_plan, project_plan_json
from ai_system.template.ooxml_parser import parse_docx
from ai_system.template.spec_builder import build_template_spec
from ai_system.validate.runner import validate_document


def _add_outline_style(document: Document, name: str, level: int, size_pt: int = 16) -> None:
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    style.font.size = Pt(size_pt)
    p_pr = style.element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def _make_course_template(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "一级标题", 0, 18)
    _add_outline_style(doc, "二级标题", 1, 14)
    doc.add_paragraph("1 实验目的", style="一级标题")
    doc.add_paragraph("1.1 背景说明", style="二级标题")
    doc.add_paragraph("学号：20260001")
    doc.add_paragraph("请在此处填写实验步骤")
    doc.add_paragraph("例如：print(\"hello\")")
    doc.add_paragraph("成稿后删除：格式要求使用宋体小四")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "示例行"
    table.cell(1, 1).text = "—"
    doc.save(str(path))
    return path


def _handwritten_ir(spec) -> PaperIR:
    sections = {}
    for slot in spec.slots:
        if slot.role == "placeholder_fill":
            sections[slot.id] = SectionContent(
                slot_id=slot.id,
                blocks=[
                    Paragraph(
                        text="本节给出实验环境、操作步骤和预期结果，内容覆盖模板占位要求并超过最小字数。"
                    )
                ],
            )
        elif slot.role == "table":
            sections[slot.id] = SectionContent(
                slot_id=slot.id,
                blocks=[
                    TableRows(
                        table_id=slot.anchor_block,
                        rows=[["项目", "结果"], ["吞吐", "120"]],
                    )
                ],
            )
    return PaperIR(work_id="work-golden", revision=1, sections=sections)


def test_infer_role_covers_template_cues():
    assert infer_role("paragraph", "1 实验目的", is_heading=True) == "heading"
    assert infer_role("paragraph", "请在此处填写实验步骤") == "placeholder_fill"
    assert infer_role("paragraph", "例如：print(\"hello\")") == "example_delete"
    assert infer_role("paragraph", "成稿后删除：格式要求使用宋体小四") == "instruction_delete"
    assert infer_role("paragraph", "学号：20260001") == "fixed_text"
    assert infer_role("table", "项目 / 结果") == "table"
    assert infer_role("paragraph", "") == "other"


def test_template_spec_labels_chinese_course_template(tmp_path: Path):
    template = _make_course_template(tmp_path / "template.docx")
    spec = build_template_spec(template, template_id=42, judge=HeuristicJudge())
    roles = {slot.role for slot in spec.slots}
    assert "heading" in roles
    assert "placeholder_fill" in roles
    assert "example_delete" in roles
    assert "instruction_delete" in roles
    assert "table" in roles
    assert "fixed_text" in roles
    headings = [slot.title for slot in spec.slots if slot.role == "heading"]
    assert headings == ["1 实验目的", "1.1 背景说明"]


def test_handwritten_ir_renders_and_validates_clean(tmp_path: Path):
    template = _make_course_template(tmp_path / "template.docx")
    spec = build_template_spec(template, template_id=42)
    ir = _handwritten_ir(spec)
    output = tmp_path / "paper.docx"
    render_docx(spec, ir, template, output)

    parsed = parse_docx(output)
    paper_text = parsed.text
    assert "1 实验目的" in paper_text
    assert "学号：20260001" in paper_text
    assert "本节给出实验环境" in paper_text
    assert "例如：print(\"hello\")" not in paper_text
    assert "成稿后删除：格式要求使用宋体小四" not in paper_text
    assert parsed.table_count == 1
    assert "吞吐" in " ".join(cell for row in parsed.blocks if row.kind == "table" for cell in sum(row.table_rows, []))

    issues = validate_document(spec, ir, output, template)
    assert [issue.model_dump() for issue in issues] == []


def test_docx_render_is_idempotent(tmp_path: Path):
    template = _make_course_template(tmp_path / "template.docx")
    spec = build_template_spec(template, template_id=7)
    ir = _handwritten_ir(spec)
    first = tmp_path / "a.docx"
    second = tmp_path / "b.docx"
    render_docx(spec, ir, template, first)
    render_docx(spec, ir, template, second)
    assert normalize_docx_bytes(first) == normalize_docx_bytes(second)


def test_markdown_render_and_validate(tmp_path: Path):
    template = _make_course_template(tmp_path / "template.docx")
    spec = build_template_spec(template, template_id=8)
    ir = _handwritten_ir(spec)
    markdown_path = tmp_path / "paper.md"
    text = render_markdown(spec, ir, markdown_path)
    assert "1 实验目的" in text
    assert "本节给出实验环境" in text
    assert "例如：print" not in text
    assert "成稿后删除" not in text
    issues = validate_document(spec, ir, markdown_path)
    assert issues == []


def test_plan_is_derived_from_slots_not_keywords(tmp_path: Path):
    template = _make_course_template(tmp_path / "template.docx")
    spec = build_template_spec(template, template_id=9)
    empty_plan = derive_plan(spec)
    writable = [slot.id for slot in spec.slots if slot.role in {"placeholder_fill", "table", "figure_slot", "caption"}]
    assert [task.slot_id for task in empty_plan.tasks] == writable
    assert all(task.status == "pending" for task in empty_plan.tasks)

    ir = _handwritten_ir(spec)
    filled = derive_plan(spec, ir)
    assert {task.status for task in filled.tasks} == {"committed"}
    projected = project_plan_json(filled, title="实验计划")
    assert projected["source"] == "template_spec"
    assert projected["stats"]["completed"] == len(filled.tasks)
    assert projected["stats"]["pending"] == 0


def test_event_emitter_dual_writes_json_block_and_file(tmp_path: Path):
    emitter = EventEmitter(workspace_dir=tmp_path, run_id="run-1", thread_id="work-1")
    event = emitter.emit_from_json_block("main_agent_start", "开始")
    assert event.event_type == "RUN_STARTED"
    assert (tmp_path / ".system" / "run_events.jsonl").exists()

    manager = StreamOutputManager(run_id="run-2", thread_id="work-2")
    asyncio.run(manager.send_json_block("plan_updated", '{"items": []}'))
    assert manager.event_emitter.events[0].event_type == "STATE_DELTA"
    assert manager.event_emitter.events[0].payload["type"] == "plan_updated"


def test_sandbox_runs_in_workdir(tmp_path: Path):
    result = run_sandbox("print('ok')\nopen('note.txt','w').write('hi')\n", tmp_path)
    assert result.returncode == 0
    assert "ok" in result.stdout
    assert "note.txt" in result.artifacts
