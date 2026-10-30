"""Independent byte trees. No actual user Python, compiler, network, or model."""
import asyncio
import base64
import hashlib
import io
import json
from pathlib import Path
import shlex
import shutil
import sys
import zipfile
from types import SimpleNamespace
import pytest
from docx import Document
from fastapi import HTTPException
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from helpers.remote_files import RemoteFiles
from ai_system.sandbox.remote_kernel import RemoteKernel
from ai_system.sandbox.settings import SandboxSettings
from ai_system.sandbox.contracts import ExecutionOutcome
from ai_system.sandbox.evidence import verify_run,freeze_inputs,seal_run
from ai_system.sandbox.artifacts import capture_local_outputs,capture_run_files
from services.file_services.workspace_fs import WorkspaceFS
from services.file_services.workspace_files import WorkspaceFileService
from services.chat_services import paper_project


@pytest.fixture(autouse=True)
def storage(tmp_path,monkeypatch):
    monkeypatch.setenv('PA_DATA_PATH',str(tmp_path/'backend-data'))
    monkeypatch.delenv('PAPERAGENT_SANDBOX_CONFIG',raising=False)


@pytest.fixture
def root(tmp_path):
    path=tmp_path/'backend-data/project-files/project';path.mkdir(parents=True)
    return path


class Cells(RemoteFiles):
    async def execute(self,code,run_id,output):
        self.calls.append(code)
        value=(self.remote/'data/input.csv').read_text()
        if code!='unchanged':
            (self.remote/'outputs/value.csv').write_text(value+'|'+code)
        await output('stdout',value+'\n')
        return ExecutionOutcome()


def test_correct_cached_inputs_and_live_cells_transfer_only_changed_files(root,tmp_path):
    async def check():
        session=Cells(root,tmp_path/'remote-host');kernel=RemoteKernel(root,session,session.settings)
        kernel.fs.write_text('data/input.csv','correct')
        await session.start();(session.remote/'data/input.csv').write_text('stale remote bytes')
        try:
            first=await kernel.execute('first')
            assert first.ok and not first.delivery_error
            assert session.uploads==['data/input.csv']
            assert (root/'outputs/value.csv').read_text()=='correct|first'
            assert not session.remote.is_relative_to(root)
            count=len(session.downloads)
            second=await kernel.execute('unchanged')
            assert second.ok and not second.files and len(session.downloads)==count
            assert session.uploads==['data/input.csv']
            kernel.fs.write_text('data/input.csv','new version')
            third=await kernel.execute('third')
            assert third.ok and session.uploads==['data/input.csv','data/input.csv']
            assert (root/'outputs/value.csv').read_text()=='new version|third'
        finally:await kernel.close()
    asyncio.run(check())


def test_rebuilt_container_reuses_confirmed_files_without_claiming_memory(root,tmp_path):
    async def check():
        first=Cells(root,tmp_path/'remote-host');kernel=RemoteKernel(root,first,first.settings)
        kernel.fs.write_text('data/input.csv','persisted')
        result=await kernel.execute('original');await kernel.close()
        assert result.ok
        replacement=Cells(root,tmp_path/'remote-host');rebuilt=RemoteKernel(root,replacement,replacement.settings)
        try:
            await rebuilt.execute('new-cell')
            assert not replacement.uploads
            assert replacement.calls==['new-cell'] and replacement.context_id!=first.context_id
            assert (replacement.remote/'data/input.csv').read_text()=='persisted'
        finally:await rebuilt.close()
    asyncio.run(check())


def test_generated_files_in_arbitrary_directories_are_delivered_without_input_changes(root,tmp_path):
    from ai_system.sandbox.transfer import output_path
    class Nested(Cells):
        async def execute(self,code,run_id,output):
            self.calls.append(code)
            for name in ('results/dispatch.csv','analysis/tables/check.json'):
                path=self.remote/name;path.parent.mkdir(parents=True,exist_ok=True);path.write_text('computed')
            (self.remote/'attachment/input.csv').write_text('changed remotely')
            return ExecutionOutcome()
    async def check():
        session=Nested(root,tmp_path/'remote-host');kernel=RemoteKernel(root,session,session.settings)
        kernel.fs.write_text('attachment/input.csv','original input')
        try:
            result=await kernel.execute('once')
            assert result.ok and not result.delivery_error
            assert (root/'results/dispatch.csv').read_text()=='computed'
            assert (root/'analysis/tables/check.json').read_text()=='computed'
            assert (root/'attachment/input.csv').read_text()=='original input'
            assert session.calls==['once']
        finally:await kernel.close()
    asyncio.run(check())
    for name in ('attachment/input.csv','data/source.csv','outputs/research/run/result.csv','paper.md'):
        with pytest.raises(ValueError):output_path(name)


def test_delivery_retry_uses_frozen_version_after_mutable_cache_changed(root,tmp_path):
    async def check():
        first=Cells(root,tmp_path/'remote-host');first.fail_download=True
        kernel=RemoteKernel(root,first,first.settings);kernel.fs.write_text('data/input.csv','data')
        result=await kernel.execute('computed-once')
        assert result.ok and result.delivery_error and first.calls==['computed-once']
        info=kernel.fs.get_run_info(result.run_id)
        assert info['status']=='success' and info['output_manifest']['files']
        assert not (root/'outputs/value.csv').exists()
        await kernel.close()
        (first.remote/'outputs/value.csv').write_text('later scratch mutation')
        second=Cells(root,tmp_path/'remote-host');rebuilt=RemoteKernel(root,second,second.settings)
        try:
            await rebuilt.retry_outputs()
            verify_run(rebuilt.fs,result.run_id)
            assert not second.calls and not second.uploads
            assert (root/'outputs/value.csv').read_text()=='data|computed-once'
            assert (second.remote/'outputs/value.csv').read_text()=='later scratch mutation'
        finally:await rebuilt.close()
    asyncio.run(check())


def test_bad_download_is_atomic_and_never_registers_evidence(root,tmp_path):
    async def check():
        session=Cells(root,tmp_path/'remote');session.corrupt_download=True
        kernel=RemoteKernel(root,session,session.settings)
        kernel.fs.write_text('data/input.csv','data');kernel.fs.write_text('outputs/value.csv','existing user file')
        result=await kernel.execute('changed')
        try:
            assert result.ok and result.delivery_error
            assert (root/'outputs/value.csv').read_text()=='existing user file'
            assert not kernel.fs.get_run_info(result.run_id).get('evidence_archived')
            assert not list(kernel.fs.run_dir(result.run_id).rglob('.download-*'))
            session.corrupt_download=False;await kernel.retry_outputs()
            assert session.calls==['changed']
            assert (root/'outputs/value.csv').read_text()=='data|changed'
        finally:await kernel.close()
    asyncio.run(check())


def test_retry_refuses_to_overwrite_a_locally_edited_working_copy(root,tmp_path):
    async def check():
        session=Cells(root,tmp_path/'remote');session.fail_download=True
        kernel=RemoteKernel(root,session,session.settings);kernel.fs.write_text('data/input.csv','data')
        result=await kernel.execute('one');await kernel.close()
        (root/'outputs/value.csv').write_text('manual change while delivery pending')
        second=Cells(root,tmp_path/'remote');rebuilt=RemoteKernel(root,second,second.settings)
        try:
            with pytest.raises(ValueError,match='拒绝'):await rebuilt.retry_outputs()
            assert (root/'outputs/value.csv').read_text()=='manual change while delivery pending'
            assert not second.calls and not rebuilt.fs.get_run_info(result.run_id).get('evidence_archived')
        finally:await rebuilt.close()
    asyncio.run(check())


@pytest.mark.parametrize('unsafe', [
    '../secret', '/etc/passwd', 'outputs/../secret', 'outputs/../../secret',
    'outputs\\secret', 'outputs/link/private', 'manifest.json', 'chat_history.json',
    '.system/research.sqlite', 'outputs/../manifest.json',
    'data/overwrite.csv', 'outputs/research/overwrite.png',
])
def test_remote_manifest_cannot_escape_backend_archive(root,tmp_path,unsafe):
    async def check():
        class Malicious(Cells):
            async def capture(self,run_id,baseline,**kwargs):
                return {'run_id':run_id,'files':[{'path':unsafe,'size':1,'sha256':'0'*64,'snapshot':'.pa-runs/'+run_id+'/files/'+unsafe}]}
        session=Malicious(root,tmp_path/'remote');kernel=RemoteKernel(root,session,session.settings)
        kernel.fs.write_text('data/input.csv','x')
        if unsafe=='outputs/link/private':
            (root/'outputs/link').symlink_to(tmp_path,target_is_directory=True)
        try:
            result=await kernel.execute('one')
            assert not kernel.fs.get_run_info(result.run_id).get('evidence_archived')
            assert not session.downloads
        finally:await kernel.close()
    asyncio.run(check())


class Compiler(RemoteFiles):
    def __init__(self,*args,**kwargs):
        super().__init__(*args,**kwargs);self.commands=[];self.reference=None;self.fail_export=False
    async def exec(self,argv,*,stdin=None,timeout=180):
        self.commands.append((list(argv), stdin))
        markdown=stdin or ''
        if argv[0]=='xelatex':
            assert (self.remote/'paper.tex').is_file() and (self.remote/'custom.sty').read_text()=='local dependency'
            build=self.remote/'.build/latex';assert build.is_dir(), 'adapter must create remote compiler output directory'
            (build/'paper.log').write_text('Output written on paper.pdf (1 page).\npass '+str(len(self.commands)))
            (build/'paper.aux').write_text('no bibliography')
            (build/'paper.pdf').write_bytes(b'mock PDF delivered by API')
        elif '--print-default-data-file' in argv:
            target=self.remote/'.build/reference.docx';target.parent.mkdir(parents=True,exist_ok=True)
            doc=Document();doc.add_paragraph('default');doc.save(target)
        elif argv[0]=='pandoc':
            reference=next(p.split('=',1)[1] for p in argv if p.startswith('--reference-doc='))
            assert reference.startswith('/workspace/')
            self.reference=(self.remote/reference.removeprefix('/workspace/')).read_bytes()
            assert markdown=='\nText' or 'Text' in markdown
            target=self.remote/'.build/export.docx';target.parent.mkdir(parents=True,exist_ok=True)
            doc=Document();doc.add_paragraph('mock export');doc.save(target)
        else:raise AssertionError('No real compiler is run')
        return 0,'mock completion',''
    async def download(self,snapshot,destination,max_bytes):
        if self.fail_export and snapshot.endswith('/.build/export.docx'):raise ConnectionError('export delivery offline')
        await super().download(snapshot,destination,max_bytes)


@pytest.fixture
def compiler_factory(tmp_path,monkeypatch):
    sessions=[]
    settings=SandboxSettings(provider='compiler-test',api_key='fake-key')
    def create(config,*,workspace_dir,cache_key=None,**kwargs):
        value=Compiler(workspace_dir,tmp_path/'remote-compiler',cache_key=cache_key,settings=config)
        sessions.append(value);return value
    from ai_system.sandbox import factory
    monkeypatch.setitem(factory._PROVIDERS,'compiler-test',create)
    monkeypatch.setattr(SandboxSettings,'from_env',classmethod(lambda cls:settings))
    return sessions


def test_latex_compiler_files_cross_api_and_completed_passes_are_not_reexecuted(root,compiler_factory):
    from ai_system.export.latex import compile_pdf
    (root/'paper.tex').write_text('\\documentclass{article}\\begin{document}Text\\end{document}')
    (root/'custom.sty').write_text('local dependency')
    (root/'data').mkdir(exist_ok=True);(root/'data/unrelated.csv').write_text('not a paper dependency')
    async def check():
        first=await compile_pdf(root)
        assert first.success and (root/'paper.pdf').read_bytes()==b'mock PDF delivered by API'
        session=compiler_factory[0]
        assert len(session.commands)==3 and 'custom.sty' in session.uploads
        assert 'data/unrelated.csv' not in session.uploads
        assert session.downloads and session.remote!=root
        (session.remote/'.build/latex/paper.pdf').write_bytes(b'changed remote scratch')
        second=await compile_pdf(root)
        assert second.success and compiler_factory[1].commands==[]
        assert (root/'paper.pdf').read_bytes()==b'mock PDF delivered by API'
        assert WorkspaceFileService().read_file('project','paper.pdf')['type']=='binary'
    asyncio.run(check())


def test_default_word_reference_roundtrip_and_failed_output_retry_across_staging(root,compiler_factory,monkeypatch):
    from ai_system.export.docx import export_docx
    # Install a one-shot output delivery fault; source and reference transfers still work.
    import ai_system.sandbox.commands as adapter
    factory=adapter.create_session
    def create(*args,**kwargs):
        session=factory(*args,**kwargs);session.fail_export=len(compiler_factory)==1;return session
    monkeypatch.setattr(adapter,'create_session',create)
    async def check():
        first=root/'.build/first';first.mkdir(parents=True);(first/'paper.md').write_text('# Title\nText')
        with pytest.raises(RuntimeError,match='delivery offline'):await export_docx(first,'# Title\nText')
        remote=compiler_factory[0]
        assert len(remote.commands)==2
        assert '.build/reference.docx' in remote.uploads, 'backend-styled reference must return through the API'
        styled=Document(io.BytesIO(remote.reference))
        assert styled.styles['Normal'].font.name=='Times New Roman'
        shutil.rmtree(first)
        second=root/'.build/second';second.mkdir();(second/'paper.md').write_text('# Title\nText')
        result=await export_docx(second,'# Title\nText')
        assert result.reference=='default' and (second/'paper.docx').is_file()
        assert compiler_factory[1].commands==[], 'retry delivery, never repeat the completed conversion'
        assert compiler_factory[1].cache_key==remote.cache_key
    asyncio.run(check())


def test_uploaded_word_template_is_available_only_after_api_upload(root,compiler_factory):
    from ai_system.export.docx import export_docx
    template=root/'template/original.docx';template.parent.mkdir();doc=Document();doc.add_paragraph('Template');doc.save(template)
    async def check():
        result=await export_docx(root,'# Title\nText',template_docx=template)
        session=compiler_factory[0]
        assert result.reference=='template' and session.reference==template.read_bytes()
        assert 'template/original.docx' in session.uploads and len(session.commands)==1
        assert (root/'paper.docx').is_file()
    asyncio.run(check())


def test_formal_outputs_ignore_modified_presentation_copy_and_are_readonly(root):
    source=WorkspaceFS(root/'experiments/attempt/workspace');source.init_structure();run=source.create_run()
    source.save_run_input(run,'print(1)');source.save_run_stdout(run,'1');source.save_run_stderr(run,'');freeze_inputs(source,run)
    source.write_text('outputs/chart.png','original figure bytes');capture_local_outputs(source,run,{})
    source.finish_run(run,'success');seal_run(source,run)
    target=WorkspaceFS(root);target.init_structure();capture_run_files(source,target)
    rel=f'outputs/research/{run}/outputs/chart.png'
    (root/'paper.tex').write_text('\\documentclass{article}\\begin{document}\\includegraphics{'+rel+'}\\end{document}')
    version=paper_project.snapshot(root)['version']
    with pytest.raises(ValueError,match='只读'):paper_project.safe_path(root,rel,writing=True)
    target.abs(rel).write_text('tampered working copy')
    service=WorkspaceFileService()
    assert base64.b64decode(service.read_file('project',rel)['content'])==b'original figure bytes'
    assert paper_project.snapshot(root)['version']==version
    with pytest.raises(HTTPException):service.write_file('project',rel,'replace evidence')
    exported=Path(service.export_workspace('project'))
    try:
        with zipfile.ZipFile(exported) as archive:assert archive.read(rel)==b'original figure bytes'
    finally:shutil.rmtree(exported.parent)


def test_host_report_can_be_listed_read_and_exported_without_execution_receipt(root):
    from helpers.research import make_store,finish
    store=make_store(root);finish(store,'Q1')
    state=store.snapshot();artifact=state['artifacts'][state['experiments']['Q1']['report_ref']]
    path=artifact['payload']['path'];fs=WorkspaceFS(root)
    frozen=fs.canonical_artifact(path)
    assert json.loads(frozen.read_text())==artifact['report']
    # A pre-fix report has only a committed DB record and a public copy.
    frozen.unlink();(root/path).write_text('tampered presentation copy')
    service=WorkspaceFileService()
    assert path in {row['path'] for row in service.list_files('project')}
    assert json.loads(service.read_file('project',path)['content'])==artifact['report']
    exported=Path(service.export_workspace('project'))
    try:
        with zipfile.ZipFile(exported) as archive:
            assert json.loads(archive.read(path))==artifact['report']
    finally:shutil.rmtree(exported.parent)
    with pytest.raises(ValueError,match='缺少归档映射'):
        fs.canonical_artifact('outputs/research/unregistered/report-0000000000000000.json')
    frozen.write_text('tampered private archive')
    with pytest.raises(ValueError,match='归档校验失败'):fs.canonical_artifact(path)


def test_download_and_range_read_the_opened_descriptor_on_this_platform(root):
    import httpx
    from fastapi import FastAPI
    source=root/'paper.pdf';raw=b'%PDF-1.4\nretained file';source.write_bytes(raw)
    service=WorkspaceFileService();app=FastAPI()
    @app.get('/download')
    async def download():
        response=service.file_response('project','paper.pdf')
        # Renaming/replacing the pathname must not change the checked open file.
        source.rename(root/'saved.pdf');source.write_bytes(b'different file')
        return response
    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
            response=await client.get('/download')
            assert response.status_code==200 and response.content==raw
            source.write_bytes(raw)
            response=await client.get('/download',headers={'Range':'bytes=0-3'})
            assert response.status_code==206 and response.content==b'%PDF'
            assert response.headers['content-range']==f'bytes 0-3/{len(raw)}'
    asyncio.run(run())


def test_provider_uses_real_file_api_adapter_with_separate_sdk_filesystem(root,tmp_path,monkeypatch):
    from opensandbox import Sandbox
    from code_interpreter import CodeInterpreter
    from ai_system.sandbox.opensandbox_provider import OpenSandboxSession
    from ai_system.sandbox.transfer import WorkspaceTransfer
    from ai_system.sandbox import transfer_worker
    disk=tmp_path/'sdk-server-volume';disk.mkdir()
    api_calls=[];mounts=[]
    def remote_path(value):
        if value.startswith('/workspace/') or value=='/workspace':return disk/value.removeprefix('/workspace').lstrip('/')
        assert value.startswith('/tmp/')
        return tmp_path/'sdk-container-tmp'/Path(value).name
    class Files:
        async def write_file(self,path,content,mode=None):
            api_calls.append(('write',path))
            value=content.read() if hasattr(content,'read') else content
            target=remote_path(path);target.parent.mkdir(parents=True,exist_ok=True)
            target.write_bytes(value.encode() if isinstance(value,str) else value)
        async def read_file(self,path,**kwargs):return remote_path(path).read_text()
        async def create_directories(self,entries):
            for entry in entries:remote_path(entry.path).mkdir(parents=True,exist_ok=True)
        async def delete_files(self,paths):
            for path in paths:remote_path(path).unlink(missing_ok=True)
        async def read_bytes_stream(self,path):
            api_calls.append(('read',path));raw=remote_path(path).read_bytes()
            async def chunks():
                for i in range(0,len(raw),3):yield raw[i:i+3]
            return chunks()
    class Commands:
        async def run(self,command,*,opts,handlers):
            argv=shlex.split(command)
            assert argv[:2]==['python3','/tmp/paperagent_transfer.py']
            prior=transfer_worker.ROOT;transfer_worker.ROOT=disk
            try:result=transfer_worker.operation(json.loads(argv[2]))
            finally:transfer_worker.ROOT=prior
            await handlers.on_stdout(SimpleNamespace(text=json.dumps(result)))
            return SimpleNamespace(error=None,exit_code=0)
    async def nothing(*args,**kwargs):pass
    async def create(*args,**kwargs):
        mounts.extend(kwargs['volumes'])
        return SimpleNamespace(id='sdk-mock',files=Files(),commands=Commands(),destroy=nothing)
    async def interpreter(*args,**kwargs):
        async def context(*args):return SimpleNamespace(id='mock-context')
        async def setup(*args,**kwargs):return SimpleNamespace(error=None,complete=True)
        return SimpleNamespace(codes=SimpleNamespace(create_context=context,run=setup),aclose=nothing)
    monkeypatch.setattr(Sandbox,'create',create);monkeypatch.setattr(CodeInterpreter,'create',interpreter)
    async def check():
        settings=SandboxSettings(api_key='fake',host_workspace_root='/server-only/cache')
        session=OpenSandboxSession(settings,workspace_dir=root)
        fs=WorkspaceFS(root);fs.init_structure();fs.write_text('data/input.csv','backend bytes')
        transfer=WorkspaceTransfer(fs,session,settings)
        try:
            await session.start();await transfer.sync_inputs()
            assert (disk/'data/input.csv').read_text()=='backend bytes'
            assert ('write','/workspace/.pa-inputs.json') not in api_calls, 'cache ledger uses atomic install'
            assert len(mounts)==1 and mounts[0].host.path=='/server-only/cache/'+session.cache_key
            assert str(root) not in mounts[0].host.path
            (disk/'outputs/result.csv').write_text('remote bytes')
            run=fs.create_run();manifest=await transfer.obtain_manifest(run,{})
            await transfer.deliver(run,manifest,{})
            assert (root/'outputs/result.csv').read_text()=='remote bytes'
            assert any(op=='read' and '/.pa-runs/' in name for op,name in api_calls)
        finally:await session.close()
    asyncio.run(check())


def test_late_figure_finalization_is_frozen_only_after_success(root,tmp_path):
    async def check():
        class Figures(Cells):
            offline=True
            async def finalize(self,run_id):
                if self.offline:raise RuntimeError('figures not saved')
                (self.remote/'outputs/figure.png').write_bytes(b'figure')
        session=Figures(root,tmp_path/'remote');kernel=RemoteKernel(root,session,session.settings)
        kernel.fs.write_text('data/input.csv','data')
        try:
            result=await kernel.execute('one')
            assert result.delivery_error and not kernel.fs.get_run_info(result.run_id).get('output_manifest')
            session.offline=False;await kernel.retry_outputs()
            verify_run(kernel.fs,result.run_id)
            assert (root/'outputs/figure.png').read_bytes()==b'figure' and session.calls==['one']
        finally:await kernel.close()
    asyncio.run(check())


def test_live_input_mutation_is_archived_without_merging_into_backend(root,tmp_path):
    async def check():
        session=Cells(root,tmp_path/'remote');kernel=RemoteKernel(root,session,session.settings)
        kernel.fs.write_text('data/input.csv','original input')
        try:
            await kernel.execute('first')
            (session.remote/'data/input.csv').write_text('live changed bytes')
            (session.remote/'data/generated.csv').write_text('live generated input')
            result=await kernel.execute('second')
            assert result.ok and not result.delivery_error
            verify_run(kernel.fs,result.run_id)
            archived=kernel.fs.run_dir(result.run_id)/'inputs/data'
            assert (archived/'input.csv').read_text()=='live changed bytes'
            assert (archived/'generated.csv').read_text()=='live generated input'
            assert (root/'data/input.csv').read_text()=='original input'
            assert not (root/'data/generated.csv').exists()
            assert session.uploads==['data/input.csv'] and session.calls==['first','second']
        finally:await kernel.close()
    asyncio.run(check())


def test_unknown_compiler_completion_is_not_replayed(root,compiler_factory,monkeypatch):
    from ai_system.export.latex import compile_pdf
    (root/'paper.tex').write_text('\\documentclass{article}\\begin{document}Text\\end{document}')
    (root/'custom.sty').write_text('local dependency')
    import ai_system.sandbox.commands as adapter
    factory=adapter.create_session
    def create(*args,**kwargs):
        session=factory(*args,**kwargs)
        async def interrupted(*args,**kwargs):
            session.commands.append('submission with unknown result')
            raise ConnectionError('completion disconnected')
        session.exec=interrupted
        return session
    monkeypatch.setattr(adapter,'create_session',create)
    async def check():
        first=await compile_pdf(root);second=await compile_pdf(root)
        assert not first.success and not second.success
        assert len(compiler_factory[0].commands)==1 and compiler_factory[1].commands==[]
        assert '不会自动重跑' in second.error
    asyncio.run(check())


def test_compile_snapshot_can_use_canonical_formal_asset_after_copy(root,compiler_factory):
    from helpers.research import make_store
    source=WorkspaceFS(root/'experiments/attempt/workspace');source.init_structure();run=source.create_run()
    source.save_run_input(run,'figure');source.save_run_stdout(run,'');source.save_run_stderr(run,'')
    source.write_text('outputs/chart.png','sealed original figure');capture_local_outputs(source,run,{})
    source.finish_run(run,'success');seal_run(source,run)
    target=WorkspaceFS(root);target.init_structure();capture_run_files(source,target)
    rel=f'outputs/research/{run}/outputs/chart.png'
    (root/'paper.tex').write_text('\\documentclass{article}\\begin{document}\\includegraphics{'+rel+'}\\end{document}')
    (root/'custom.sty').write_text('local dependency')
    target.abs(rel).write_text('tampered presentation')
    store=make_store(root)
    async def check():
        result=await paper_project.compile_project(store)
        assert result['compile']['success']
        session=compiler_factory[0]
        assert (session.remote/rel).read_text()=='sealed original figure'
        assert rel in session.uploads and (root/'paper.pdf').is_file()
        from ai_system.export.docx import export_docx
        await export_docx(root,'# Title\nText\n![plot]('+rel+')')
        assert (compiler_factory[1].remote/rel).read_text()=='sealed original figure'
        assert target.abs(rel).read_text()=='tampered presentation'
    asyncio.run(check())


def test_word_uploads_only_selected_images_and_reference(root,compiler_factory):
    from ai_system.export.docx import export_docx
    image=root/'figures/chart.png';image.parent.mkdir();image.write_bytes(b'selected image')
    unrelated=root/'data/raw.csv';unrelated.parent.mkdir();unrelated.write_text('not a Word dependency')
    async def check():
        result=await export_docx(root,'# Title\nText\n![plot](figures/chart.png)')
        session=compiler_factory[0]
        assert result.docx_path=='paper.docx'
        assert (session.remote/'figures/chart.png').read_bytes()==image.read_bytes()
        assert set(session.uploads)=={'figures/chart.png','.build/reference.docx'}
        assert not (session.remote/'data/raw.csv').exists()
    asyncio.run(check())
