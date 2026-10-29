"""Regressions from the real microgrid run: IDs, state ownership and checks."""
import asyncio
import json
from pathlib import Path
import pytest
from langchain_core.messages import AIMessage,ToolMessage
from ai_system.modules.registry import invoke,capabilities
from ai_system.research.runtime import execution_scope
from test_context_sessions import context


@pytest.mark.usefixtures('offline_kernel')
def test_new_questions_and_public_code_need_no_experiment_or_acceptance(tmp_path):
    ctx=context(tmp_path)
    async def run():
        try:
            public=await invoke(ctx,'research.run_python',code="print('Inspecting original data', 144)")
            first=await invoke(ctx,'research.create_question',question='单日调度')
            second=await invoke(ctx,'research.create_question',question='全年调度')
            assert (first['question_id'],second['question_id'])==('Q1','Q2')
            with pytest.raises(ValueError,match='编号不存在'):
                await invoke(ctx,'research.run_python',question_id='invented',code='print(3)')
            owned=await invoke(ctx,'research.run_python',question_id='Q2',code='print(365)')
            assert public['question_id'] is None and owned['question_id']=='Q2'
            state=ctx.research.snapshot()
            assert not state['experiments'] and not state['executions'] and not state['artifacts']
            assert len(ctx.fs._manifest['runs'])==2
            assert all('acceptance' not in r for r in ctx.fs._manifest['runs'].values())
            with pytest.raises(ValueError,match='编号不存在'):
                await invoke(ctx,'research.edit_question',question_id='bad',base_version=1,question='改题')
        finally:await ctx.kernels.close()
    asyncio.run(run())
    names={c.name for c in capabilities().model_capabilities()}
    assert {'research.create_question','research.edit_question'}<=names
    assert not {'research.update_question','research.continue_experiment','research.check_interpretation'}&names


def test_retained_tool_history_keeps_executable_code_and_reasoning(tmp_path):
    from ai_system.graph.compaction import ContextMiddleware
    from research_tools_helpers import ResearchToolsModel,tool
    ctx=context(tmp_path)
    messages=[]
    for i in range(4):
        path=f'code/cell_{i}.py';code='data = '+repr('x'*4000)
        ctx.fs.write_text(path,code);ctx.fs.register_file(path,kind='code',visibility='user')
        messages += [tool('research_run_python',{'code':code},str(i)),
                     ToolMessage(content=json.dumps({'input_path':path,'run_id':f'run_{i}'}),tool_call_id=str(i),name='research_run_python')]
        messages[-2].additional_kwargs['reasoning_content']='offline reasoning'
    class Request:
        tools=[]
        def __init__(self,messages):self.messages=messages
        def override(self,**kw):return Request(kw['messages'])
    sent=[]
    async def handler(r):sent.extend(r.messages)
    model=ResearchToolsModel(respond=lambda m:AIMessage(content='notes'))
    asyncio.run(ContextMiddleware(ctx,model,'reduce-code').awrap_model_call(Request(messages),handler))
    assert sent[0].tool_calls[0]['args']['code']==ctx.fs.read_text('code/cell_0.py')
    assert all(m.additional_kwargs['reasoning_content']=='offline reasoning' for m in sent if isinstance(m,AIMessage))
    assert messages[0].tool_calls[0]['args']['code']==ctx.fs.read_text('code/cell_0.py')
    assert sent[-2].tool_calls[0]['args']['code']==messages[-2].tool_calls[0]['args']['code']


def test_microgrid_grader_uses_independent_energy_balance_even_at_correct_cost(tmp_path):
    from evals.microgrid import grade_result1
    import openpyxl
    source=openpyxl.Workbook();sheet=source.active
    sheet.append(['时间','电价','负载','光伏'])
    for i in range(144):sheet.append([i,1,100,0])
    source.save(tmp_path/'source.xlsx')
    result=openpyxl.Workbook();plan=result.active;plan.title='计划购电量'
    plan.append(['时间段','购电量'])
    for i in range(144):plan.append([i,100/6])
    storage=result.create_sheet('充放电量');storage.append(['时间段','充电量','放电量','时刻','储电量'])
    for i in range(6):storage.append([str(i),0,0,'0:00' if i==0 else '24:00' if i==1 else None,6000 if i<2 else None])
    result.save(tmp_path/'result.xlsx')
    assert grade_result1(tmp_path/'source.xlsx',tmp_path/'result.xlsx')['passed']
    storage['B2']=100  # A plausible looking but impossible charge with no SOC change.
    result.save(tmp_path/'result.xlsx')
    checked=grade_result1(tmp_path/'source.xlsx',tmp_path/'result.xlsx')
    assert checked['checks']['objective_within_2_yuan']
    assert not checked['checks']['independent_daily_energy_balance'] and not checked['passed']
