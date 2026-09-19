import asyncio
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from ai_system.eval.harness import EvalCase, run_eval_case, summarize
from ai_system.graph.checkpoint import load_checkpoint
from ai_system.graph.context import bind_emitter
from ai_system.graph.nodes.context import load_context
from ai_system.graph.nodes.gather import gather_batch, needed_gather_slots
from ai_system.graph.runner import run_pipeline
from ai_system.graph.state import PaperState
from ai_system.judge.base import Answer
from ai_system.runtime.pipeline import pipeline_version
from ai_system.schemas.paper_ir import PaperIR
from ai_system.schemas.template_spec import Block, Slot, TemplateSpec
from ai_system.template.ooxml_parser import parse_docx
from ai_system.template.spec_builder import build_template_spec
from ai_system.validate.runner import validate_document


def _add_outline_style(document: Document, name: str, level: int, size_pt: int = 16) -> None:
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    style.font.size = Pt(size_pt)
    p_pr = style.element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def _course_template(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "一级标题", 0, 18)
    _add_outline_style(doc, "二级标题", 1, 14)
    doc.add_paragraph("1 实验目的", style="一级标题")
    doc.add_paragraph("1.1 背景说明", style="二级标题")
    doc.add_paragraph("学号：20260001")
    doc.add_paragraph("请在此处填写实验步骤")
    doc.add_paragraph("例如：print(\"hello\")")
    doc.add_paragraph("成稿后删除：格式要求使用宋体小四")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "示例行"
    table.cell(1, 1).text = "—"
    doc.save(str(path))
    return path


def test_pipeline_defaults_to_v2(monkeypatch):
    monkeypatch.delenv("AI_PIPELINE", raising=False)
    assert pipeline_version() == "v2"
    monkeypatch.setenv("AI_PIPELINE", "legacy")
    assert pipeline_version() == "legacy"


def test_graph_markdown_end_to_end(tmp_path: Path):
    workspace = tmp_path / "md-work"
    workspace.mkdir()
    summary = asyncio.run(
        run_pipeline(
            work_id="md-1",
            workspace_dir=str(workspace),
            user_message="写一份关于数据库实验的课程报告",
            output_mode="markdown",
        )
    )
    assert (workspace / "paper.md").exists()
    assert (workspace / "plan.json").exists()
    assert (workspace / ".system" / "paper_ir.json").exists()
    assert (workspace / ".system" / "checkpoint.json").exists()
    text = (workspace / "paper.md").read_text(encoding="utf-8")
    assert "摘要" in text
    assert "引言" in text
    plan = json.loads((workspace / "plan.json").read_text(encoding="utf-8"))
    assert plan["source"] == "template_spec"
    assert plan["stats"]["completed"] >= 1
    assert "完成" in summary or "写入" in summary
    assert not (workspace / "outputs" / "note.txt").exists()
    assert (workspace / ".system" / "run_events.jsonl").exists()


def test_graph_word_end_to_end_validates(tmp_path: Path):
    workspace = tmp_path / "word-work"
    workspace.mkdir()
    system = workspace / ".system"
    system.mkdir()
    template = _course_template(system / "_template_original.docx")
    import shutil

    shutil.copyfile(template, workspace / "paper.docx")
    asyncio.run(
        run_pipeline(
            work_id="word-1",
            workspace_dir=str(workspace),
            user_message="按模板完成实验报告",
            template_id=42,
            output_mode="word",
        )
    )
    paper = workspace / "paper.docx"
    assert paper.exists()
    parsed = parse_docx(paper)
    assert "1 实验目的" in parsed.text
    assert "学号：20260001" in parsed.text
    assert "例如：print(\"hello\")" not in parsed.text
    spec = build_template_spec(template, 42)
    ir_data = json.loads((workspace / ".system" / "paper_ir.json").read_text(encoding="utf-8"))
    from ai_system.schemas.paper_ir import PaperIR

    ir = PaperIR.model_validate(ir_data)
    issues = validate_document(spec, ir, paper, template)
    assert issues == []


def test_graph_edit_rewrites_only_nth_section(tmp_path: Path):
    workspace = tmp_path / "edit-work"
    workspace.mkdir()
    asyncio.run(
        run_pipeline(
            work_id="edit-1",
            workspace_dir=str(workspace),
            user_message="写一篇方法论文",
            output_mode="markdown",
        )
    )
    first = json.loads((workspace / ".system" / "paper_ir.json").read_text(encoding="utf-8"))
    first_ids = set(first["sections"])
    first_rev = {key: section["revision"] for key, section in first["sections"].items()}
    asyncio.run(
        run_pipeline(
            work_id="edit-1",
            workspace_dir=str(workspace),
            user_message="修改第2节，补充更多方法细节",
            output_mode="markdown",
        )
    )
    second = json.loads((workspace / ".system" / "paper_ir.json").read_text(encoding="utf-8"))
    assert set(second["sections"]) == first_ids
    bumped = [key for key, section in second["sections"].items() if section["revision"] > first_rev[key]]
    assert len(bumped) == 1


def test_load_context_reads_chat_history(tmp_path: Path):
    (tmp_path / "chat_history.json").write_text(
        json.dumps(
            {
                "messages": [
                    {"role": "user", "content": "先写数据库实验的摘要"},
                    {"role": "assistant", "content": "已写摘要"},
                    {"role": "user", "content": "现在写到哪一步了？"},
                ]
            }
        ),
        encoding="utf-8",
    )
    state = load_context(
        PaperState(work_id="hist-1", workspace_dir=str(tmp_path), user_message="现在写到哪一步了？")
    )
    assert any("数据库实验" in turn["content"] for turn in state.history)


def test_graph_question_uses_chat_history(tmp_path: Path):
    workspace = tmp_path / "ask-hist"
    workspace.mkdir()
    (workspace / "chat_history.json").write_text(
        json.dumps(
            {
                "messages": [
                    {"role": "user", "content": "先写数据库实验的摘要"},
                    {"role": "assistant", "content": "已写摘要"},
                ]
            }
        ),
        encoding="utf-8",
    )
    summary = asyncio.run(
        run_pipeline(
            work_id="ask-hist",
            workspace_dir=str(workspace),
            user_message="现在写到哪一步了？",
            output_mode="markdown",
        )
    )
    assert "数据库实验" in summary
    assert not (workspace / "paper.md").exists()


class _RejectJudge:
    def ask(self, state, questions):
        return {"substantive": Answer(value=False, confidence=0.99)}


def test_failed_drafts_are_not_committed(tmp_path: Path):
    workspace = tmp_path / "reject-work"
    workspace.mkdir()
    asyncio.run(
        run_pipeline(
            work_id="reject-1",
            workspace_dir=str(workspace),
            user_message="写实验报告",
            output_mode="markdown",
            judge=_RejectJudge(),
            max_repair_rounds=0,
        )
    )
    ir_path = workspace / ".system" / "paper_ir.json"
    if ir_path.exists():
        ir = json.loads(ir_path.read_text(encoding="utf-8"))
        assert ir.get("sections") == {}
    else:
        assert not ir_path.exists()


def test_gather_only_targets_figure_slots(tmp_path: Path):
    spec = TemplateSpec(
        template_id=1,
        version=1,
        blocks=[
            Block(id="P001", kind="paragraph", text="请填写"),
            Block(id="G001", kind="figure", text="图1"),
        ],
        slots=[
            Slot(
                id="slot.body",
                role="placeholder_fill",
                anchor_block="P001",
                expects=["paragraph", "code"],
                title="正文",
            ),
            Slot(id="slot.fig", role="figure_slot", anchor_block="G001", title="图1"),
        ],
    )
    state = PaperState(
        work_id="g1",
        workspace_dir=str(tmp_path),
        spec=spec,
        ir=PaperIR(work_id="g1"),
    )
    assert [slot.id for slot in needed_gather_slots(state, ["slot.body", "slot.fig"])] == ["slot.fig"]
    seen: dict[str, list[str]] = {}

    def gather_fn(_state, needed):
        seen["ids"] = [slot.id for slot in needed]

    asyncio.run(gather_batch(state, ["slot.body", "slot.fig"], gather_fn=gather_fn))
    assert seen["ids"] == ["slot.fig"]
    asyncio.run(gather_batch(state, ["slot.body"], gather_fn=gather_fn))
    assert seen["ids"] == ["slot.fig"]


def test_bind_emitter_attaches_workspace(tmp_path: Path):
    class Stream:
        event_emitter = None

    stream = Stream()
    emitter = bind_emitter(stream, str(tmp_path), "run-1", "work-1")
    emitter.emit("RUN_STARTED", {"ok": True})
    assert (tmp_path / ".system" / "run_events.jsonl").exists()
    assert stream.event_emitter is emitter


def test_graph_question_does_not_write_paper(tmp_path: Path):
    workspace = tmp_path / "ask-work"
    workspace.mkdir()
    summary = asyncio.run(
        run_pipeline(
            work_id="ask-1",
            workspace_dir=str(workspace),
            user_message="现在写到哪一步了？",
            output_mode="markdown",
        )
    )
    assert not (workspace / "paper.md").exists()
    assert "计划" in summary or "需求" in summary


def test_graph_checkpoint_survives_reload(tmp_path: Path):
    workspace = tmp_path / "ckpt-work"
    workspace.mkdir()
    asyncio.run(
        run_pipeline(
            work_id="ckpt-1",
            workspace_dir=str(workspace),
            user_message="写结论部分",
            output_mode="markdown",
        )
    )
    loaded = load_checkpoint(workspace)
    assert loaded is not None
    assert loaded.work_id == "ckpt-1"
    assert loaded.ir is not None
    assert loaded.ir.sections


def test_eval_harness_marks_clean_markdown_case(tmp_path: Path):
    workspace = tmp_path / "eval-work"
    workspace.mkdir()
    asyncio.run(
        run_pipeline(
            work_id="eval-1",
            workspace_dir=str(workspace),
            user_message="写实验报告",
            output_mode="markdown",
        )
    )
    from ai_system.schemas.paper_ir import PaperIR
    from ai_system.schemas.template_spec import TemplateSpec

    spec = TemplateSpec.model_validate(json.loads((workspace / ".system" / "template_spec.json").read_text(encoding="utf-8")))
    ir = PaperIR.model_validate(json.loads((workspace / ".system" / "paper_ir.json").read_text(encoding="utf-8")))
    result = run_eval_case(EvalCase(name="md", spec=spec, ir=ir, output_mode="markdown"), tmp_path)
    summary = summarize([result])
    assert result.issue_count == 0
    assert summary["clean"] == 1
    assert summary["skeleton_ok_rate"] == 1.0
