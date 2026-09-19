import asyncio
import json
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

from ai_system.graph.checkpoint import load_checkpoint, save_checkpoint
from ai_system.graph.context import RunContext
from ai_system.graph.nodes.answer import answer
from ai_system.graph.nodes.draft import DraftError, draft_slot, heuristic_draft
from ai_system.graph.nodes.gather import promote_new_run_artifacts, register_output_artifacts
from ai_system.graph.nodes.plan import plan_node
from ai_system.graph.nodes.render_node import render_node
from ai_system.graph.nodes.spec import needs_figure, synthetic_markdown_spec
from ai_system.graph.nodes.triage import triage_issues
from ai_system.graph.runner import run_pipeline
from ai_system.graph.state import PaperState
from ai_system.render.docx_renderer import render_docx
from ai_system.render.markdown_renderer import render_markdown
from ai_system.schemas.intent import EditIntent
from ai_system.schemas.paper_ir import FigureRef, ListBlock, PaperIR, Paragraph, SectionContent, TableRows
from ai_system.schemas.template_spec import Slot
from ai_system.template.ooxml_parser import parse_docx
from ai_system.template.spec_builder import build_template_spec
from ai_system.validate.content import content_issues
from ai_system.validate.issues import ValidationIssue


class _Reply:
    def __init__(self, content: str):
        self.content = content


class _ScriptedLLM:
    """Returns scripted replies in order and records every prompt it receives."""

    def __init__(self, replies):
        self.replies = list(replies)
        self.prompts: list[str] = []

    async def ainvoke(self, messages):
        self.prompts.append(messages[0].content)
        reply = self.replies.pop(0) if len(self.replies) > 1 else self.replies[0]
        if isinstance(reply, Exception):
            raise reply
        return _Reply(reply)


def _good_draft(text: str) -> str:
    return json.dumps({"slot_id": "x", "blocks": [{"type": "paragraph", "text": text}]}, ensure_ascii=False)


LONG_TEXT = "本节系统说明实验的目的、环境配置、操作步骤与可复核的结果，并对比不同方案给出结论。" * 2


def _add_outline_style(document: Document, name: str, level: int, size_pt: int = 16) -> None:
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    style.font.size = Pt(size_pt)
    p_pr = style.element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def _two_section_template(path: Path) -> Path:
    doc = Document()
    _add_outline_style(doc, "一级标题", 0, 18)
    doc.add_paragraph("1 实验目的", style="一级标题")
    doc.add_paragraph("请在此处填写")
    doc.add_paragraph("2 实验步骤", style="一级标题")
    doc.add_paragraph("请在此处填写")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "用例"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "示例行"
    table.cell(1, 1).text = "—"
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# draft / revise
# ---------------------------------------------------------------------------


def test_llm_draft_retries_on_bad_json_with_error_feedback(tmp_path: Path):
    llm = _ScriptedLLM(["这不是 JSON", _good_draft(LONG_TEXT)])
    asyncio.run(
        run_pipeline(work_id="retry-1", workspace_dir=str(tmp_path), user_message="写实验报告", llm=llm)
    )
    ir = json.loads((tmp_path / ".system" / "paper_ir.json").read_text(encoding="utf-8"))
    assert ir["sections"], "valid second reply should be committed"
    assert any("无法解析" in prompt for prompt in llm.prompts)
    assert all(section["provenance"]["model"] != "heuristic" for section in ir["sections"].values())


def test_revise_calls_llm_with_judge_feedback_instead_of_filler(tmp_path: Path):
    spec = synthetic_markdown_spec(0, ["摘要"])
    slot = next(s for s in spec.slots if s.role == "placeholder_fill")
    llm = _ScriptedLLM([_good_draft("太短"), _good_draft(LONG_TEXT)])
    state = PaperState(work_id="rev-1", workspace_dir=str(tmp_path), spec=spec, ir=PaperIR(work_id="rev-1"))
    ctx = RunContext(llm=llm)

    from ai_system.graph.runner import _draft_and_judge

    draft, judgement = asyncio.run(_draft_and_judge(slot, state, ctx))
    assert judgement.passed
    assert draft.blocks[0].text == LONG_TEXT
    assert len(llm.prompts) == 2
    assert "内容不够充实" in llm.prompts[1]
    assert draft.provenance.model != "heuristic"


def test_llm_failure_leaves_slot_pending_instead_of_committing_filler(tmp_path: Path):
    llm = _ScriptedLLM([RuntimeError("provider down")])
    summary = asyncio.run(
        run_pipeline(work_id="fail-1", workspace_dir=str(tmp_path), user_message="写实验报告", llm=llm, max_repair_rounds=0)
    )
    ir_path = tmp_path / ".system" / "paper_ir.json"
    assert not ir_path.exists() or json.loads(ir_path.read_text(encoding="utf-8"))["sections"] == {}
    plan = json.loads((tmp_path / "plan.json").read_text(encoding="utf-8"))
    assert plan["stats"]["completed"] == 0
    assert plan["stats"]["pending"] == plan["stats"]["total"]
    assert "未能生成合格草稿" in summary


def test_heuristic_figure_draft_requires_artifact(tmp_path: Path):
    slot = Slot(id="slot.fig", role="figure_slot", anchor_block="G000", title="结果图")
    state = PaperState(work_id="fig-1", workspace_dir=str(tmp_path), ir=PaperIR(work_id="fig-1"))
    try:
        heuristic_draft(slot, state)
        raise AssertionError("expected DraftError")
    except DraftError:
        pass
    (tmp_path / "outputs").mkdir()
    Image.new("RGB", (8, 8), (10, 20, 30)).save(tmp_path / "outputs" / "plot.png")
    register_output_artifacts(state)
    draft = asyncio.run(draft_slot(slot, state, RunContext()))
    assert isinstance(draft.blocks[0], FigureRef)
    assert draft.blocks[0].artifact_id == "plot"


# ---------------------------------------------------------------------------
# plan / triage
# ---------------------------------------------------------------------------


def test_plan_edit_targets_only_and_keeps_real_statuses(tmp_path: Path):
    spec = synthetic_markdown_spec(0, ["摘要", "引言", "方法"])
    bodies = [s.id for s in spec.slots if s.role == "placeholder_fill"]
    ir = PaperIR(work_id="p1", sections={bodies[0]: SectionContent(slot_id=bodies[0], blocks=[Paragraph(text=LONG_TEXT)])})
    state = PaperState(work_id="p1", spec=spec, ir=ir, intent=EditIntent(kind="edit", target_slots=[bodies[1]]))
    state = plan_node(state)
    statuses = {task.slot_id: task.status for task in state.plan.tasks}
    assert statuses[bodies[0]] == "committed"
    assert statuses[bodies[1]] == "pending"
    assert statuses[bodies[2]] == "pending", "untouched unwritten slots must not be faked as committed"
    assert state.plan.batches == [[bodies[1]]]


def test_triage_only_retries_content_fixable_slot_issues():
    spec = synthetic_markdown_spec(0, ["摘要"])
    body = next(s.id for s in spec.slots if s.role == "placeholder_fill")
    state = PaperState(work_id="t1", spec=spec, issues=[ValidationIssue(code="skeleton_mismatch", detail="骨架")])
    state = triage_issues(state)
    assert state.awaiting_confirmation is True
    assert state.repair_round == 0

    state = PaperState(
        work_id="t2",
        spec=spec,
        issues=[
            ValidationIssue(code="placeholder_too_short", slot_id=body),
            ValidationIssue(code="fingerprint_error", detail="x", severity="warning"),
        ],
    )
    state = triage_issues(state)
    assert state.awaiting_confirmation is False
    assert state.intent.target_slots == [body]
    assert state.repair_round == 1

    state = PaperState(work_id="t3", spec=spec, issues=[ValidationIssue(code="figure_not_embedded", severity="warning")])
    state = triage_issues(state)
    assert state.awaiting_confirmation is False
    assert state.repair_round == 0


def test_repair_budget_does_not_loop_on_warnings_only(tmp_path: Path):
    spec = synthetic_markdown_spec(0, ["摘要"])
    (tmp_path / ".system").mkdir()
    (tmp_path / ".system" / "template_spec.json").write_text(spec.model_dump_json(), encoding="utf-8")
    summary = asyncio.run(
        run_pipeline(work_id="warn-1", workspace_dir=str(tmp_path), user_message="写摘要", output_mode="markdown")
    )
    assert "完成" in summary or "写入" in summary


def test_confirm_clears_pending_state_in_checkpoint(tmp_path: Path):
    asyncio.run(run_pipeline(work_id="cf-1", workspace_dir=str(tmp_path), user_message="写实验报告"))
    state = load_checkpoint(tmp_path)
    state.awaiting_confirmation = True
    save_checkpoint(state)
    asyncio.run(run_pipeline(work_id="cf-1", workspace_dir=str(tmp_path), user_message="确认采用当前稿"))
    state = load_checkpoint(tmp_path)
    assert state.awaiting_confirmation is False
    assert state.pending_confirmation is False
    events = [json.loads(line) for line in (tmp_path / ".system" / "run_events.jsonl").read_text(encoding="utf-8").splitlines()]
    assert any(event["payload"].get("type") == "confirmation_resolved" for event in events)


# ---------------------------------------------------------------------------
# renderer
# ---------------------------------------------------------------------------


def test_docx_render_replaces_placeholders_by_anchor_and_grows_tables(tmp_path: Path):
    template = _two_section_template(tmp_path / "template.docx")
    spec = build_template_spec(template, 1)
    bodies = [s for s in spec.slots if s.role == "placeholder_fill"]
    table_slot = next(s for s in spec.slots if s.role == "table")
    assert len(bodies) == 2
    ir = PaperIR(
        work_id="r1",
        sections={
            bodies[0].id: SectionContent(slot_id=bodies[0].id, blocks=[Paragraph(text="第一节正文：目的说明。" * 3)]),
            bodies[1].id: SectionContent(
                slot_id=bodies[1].id,
                blocks=[
                    Paragraph(text="第二节正文：步骤说明。" * 3),
                    ListBlock(ordered=True, items=["安装环境", "运行脚本", "记录结果"]),
                ],
            ),
            table_slot.id: SectionContent(
                slot_id=table_slot.id,
                blocks=[TableRows(table_id=table_slot.anchor_block, rows=[["用例", "结果"], ["登录", "通过"], ["注册", "通过"], ["登出", "通过"]])],
            ),
        },
    )
    output = tmp_path / "paper.docx"
    render_docx(spec, ir, template, output)
    parsed = parse_docx(output)
    texts = [item["text"] for item in parsed.paragraphs]
    assert "请在此处填写" not in texts, "placeholder prompts must be replaced, not left in place"
    assert texts.index("1 实验目的") < texts.index("第一节正文：目的说明。" * 3) < texts.index("2 实验步骤")
    assert texts.index("2 实验步骤") < texts.index("第二节正文：步骤说明。" * 3)
    assert "1. 安装环境" in texts and "3. 记录结果" in texts, "list items render as separate paragraphs"
    table = next(block for block in parsed.blocks if block.kind == "table")
    assert table.table_rows == [["用例", "结果"], ["登录", "通过"], ["注册", "通过"], ["登出", "通过"]]
    assert "示例行" not in parsed.text
    assert content_issues(spec, ir, output) == []


def test_placeholder_left_is_reported_when_prompt_survives(tmp_path: Path):
    template = _two_section_template(tmp_path / "template.docx")
    spec = build_template_spec(template, 1)
    body = next(s for s in spec.slots if s.role == "placeholder_fill")
    ir = PaperIR(work_id="pl", sections={body.id: SectionContent(slot_id=body.id, blocks=[Paragraph(text=LONG_TEXT)])})
    # Simulate a renderer that appended instead of replacing: prompt still in the doc.
    issues = content_issues(spec, ir, template)
    assert any(issue.code == "placeholder_left" and issue.slot_id == body.id for issue in issues)


def test_render_node_snapshots_template_and_rerender_does_not_duplicate(tmp_path: Path):
    template = _two_section_template(tmp_path / "paper.docx")
    spec = build_template_spec(template, 1)
    body = next(s for s in spec.slots if s.role == "placeholder_fill")
    ir = PaperIR(work_id="snap", sections={body.id: SectionContent(slot_id=body.id, blocks=[Paragraph(text=LONG_TEXT)])})
    state = PaperState(work_id="snap", workspace_dir=str(tmp_path), output_mode="word", spec=spec, ir=ir)
    render_node(state)
    assert (tmp_path / ".system" / "_template_original.docx").exists()
    first = parse_docx(tmp_path / "paper.docx").text
    render_node(state)
    second = parse_docx(tmp_path / "paper.docx").text
    assert first == second
    assert first.count(LONG_TEXT) == 1


def test_markdown_headings_follow_outline_without_fake_title():
    spec = synthetic_markdown_spec(0, ["摘要", "引言"])
    body = next(s.id for s in spec.slots if s.role == "placeholder_fill")
    ir = PaperIR(work_id="md", sections={body: SectionContent(slot_id=body, blocks=[Paragraph(text=LONG_TEXT)])})
    text = render_markdown(spec, ir)
    assert text.startswith("# 摘要\n")
    assert text.count("# 摘要") == 1
    assert "\n# 引言\n" in text


# ---------------------------------------------------------------------------
# gather / figure slot
# ---------------------------------------------------------------------------


def test_gather_promotes_new_run_artifacts_to_outputs(tmp_path: Path):
    run_dir = tmp_path / "runs" / "run_20260919_000000_abc123" / "artifacts"
    run_dir.mkdir(parents=True)
    Image.new("RGB", (8, 8), (1, 2, 3)).save(run_dir / "plot_1.png")
    (run_dir / "notes.bin").write_bytes(b"x")
    promoted = promote_new_run_artifacts(str(tmp_path), known_runs=set())
    assert promoted == ["plot_1.png"]
    assert (tmp_path / "outputs" / "plot_1.png").exists()
    state = PaperState(work_id="g", workspace_dir=str(tmp_path), ir=PaperIR(work_id="g"))
    register_output_artifacts(state)
    assert state.ir.artifacts["plot_1"].mime == "image/png"
    assert promote_new_run_artifacts(str(tmp_path), known_runs={"run_20260919_000000_abc123"}) == []


def test_synthetic_spec_adds_figure_slot_for_computational_requests():
    assert needs_figure("用蒙特卡洛模拟估计圆周率并画出收敛曲线") is True
    assert needs_figure("写一篇关于数据库的课程报告") is False
    spec = synthetic_markdown_spec(0, ["摘要", "方法", "结果", "结论"], with_figure=True)
    figure = [s for s in spec.slots if s.role == "figure_slot"]
    assert len(figure) == 1
    assert figure[0].section_path == ["结果"]
    assert not [s for s in synthetic_markdown_spec(0, ["摘要"]).slots if s.role == "figure_slot"]


def test_figure_slot_pipeline_embeds_gathered_artifact(tmp_path: Path):
    def gather_fn(state, needed):
        outputs = Path(state.workspace_dir) / "outputs"
        outputs.mkdir(exist_ok=True)
        Image.new("RGB", (8, 8), (9, 9, 9)).save(outputs / "convergence.png")

    asyncio.run(
        run_pipeline(
            work_id="mc-1",
            workspace_dir=str(tmp_path),
            user_message="用蒙特卡洛模拟估计圆周率并画出收敛曲线",
            gather_fn=gather_fn,
        )
    )
    text = (tmp_path / "paper.md").read_text(encoding="utf-8")
    assert "![" in text and "outputs/convergence.png" in text
    plan = json.loads((tmp_path / "plan.json").read_text(encoding="utf-8"))
    assert plan["stats"]["pending"] == 0


# ---------------------------------------------------------------------------
# answer streaming
# ---------------------------------------------------------------------------


def test_answer_streams_tokens_through_content_channel(tmp_path: Path):
    class StreamingLLM:
        async def astream(self, messages):
            for token in ["当前", "进度", "正常"]:
                yield _Reply(token)

        async def ainvoke(self, messages):
            raise AssertionError("astream should be preferred")

    class Stream:
        def __init__(self):
            self.chunks = []

        async def print_stream(self, delta):
            self.chunks.append(delta)

    stream = Stream()
    state = PaperState(work_id="a1", workspace_dir=str(tmp_path), user_message="进度如何？")
    state = asyncio.run(answer(state, llm=StreamingLLM(), ctx=RunContext(stream=stream)))
    assert stream.chunks == ["当前", "进度", "正常"]
    assert state.summary == "当前进度正常"
    assert state.streamed_answer is True
