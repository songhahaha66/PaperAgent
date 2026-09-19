from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from ai_system.template.ooxml_parser import parse_docx


def _add_outline_style(document: Document, name: str, level: int, size_pt: int = 16) -> None:
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    style.font.size = Pt(size_pt)
    p_pr = style.element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def _heading_texts(path: Path) -> list[str]:
    parsed = parse_docx(path)
    return [text for _, text in parsed.headings]


def test_course_report_custom_heading_styles(tmp_path: Path):
    path = tmp_path / "course-report.docx"
    doc = Document()
    _add_outline_style(doc, "一级标题", 0, 18)
    _add_outline_style(doc, "二级标题", 1, 14)
    doc.add_paragraph("1 实验目的", style="一级标题")
    doc.add_paragraph("1.1 背景说明", style="二级标题")
    doc.add_paragraph("请在此处填写实验步骤")
    doc.save(str(path))

    headings = _heading_texts(path)
    assert headings == ["1 实验目的", "1.1 背景说明"]
    parsed = parse_docx(path)
    heading_blocks = [block for block in parsed.blocks if block.is_heading]
    assert [block.outline_level for block in heading_blocks] == [0, 1]
    assert all(block.style in {"一级标题", "二级标题"} for block in heading_blocks)


def test_thesis_chapter_custom_heading_styles(tmp_path: Path):
    path = tmp_path / "thesis.docx"
    doc = Document()
    _add_outline_style(doc, "章标题", 0, 18)
    _add_outline_style(doc, "节标题", 1, 14)
    doc.add_paragraph("第一章 绪论", style="章标题")
    doc.add_paragraph("1.1 研究背景", style="节标题")
    doc.add_paragraph("学号：20260001")
    doc.save(str(path))

    assert _heading_texts(path) == ["第一章 绪论", "1.1 研究背景"]


def test_journal_numbered_title_styles_without_heading_prefix(tmp_path: Path):
    path = tmp_path / "journal.docx"
    doc = Document()
    _add_outline_style(doc, "标题 1", 0, 16)
    _add_outline_style(doc, "标题 2", 1, 14)
    doc.add_paragraph("Abstract", style="标题 1")
    doc.add_paragraph("Introduction", style="标题 2")
    doc.add_paragraph("图1 实验结果")
    doc.save(str(path))

    parsed = parse_docx(path)
    headings = [text for _, text in parsed.headings]
    assert headings == ["Abstract", "Introduction"]
    assert all(not style.startswith("heading") for style, _ in parsed.headings)
    assert all(block.style.startswith("标题") for block in parsed.blocks if block.is_heading)
