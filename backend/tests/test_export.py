"""论文导出（LaTeX / Word）与模板章节结构"""
from services.file_services.workspace_layout import workspace_path

import asyncio
import shutil
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ai_system.export.docx import export_docx
from ai_system.graph import PaperGraphAgent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from helpers.graph import TASK, FakeModel, _make_workspace

needs_pandoc = pytest.mark.skipif(not shutil.which("pandoc"), reason="没有 pandoc")
needs_xelatex = pytest.mark.skipif(not shutil.which("xelatex"), reason="没有 xelatex")

PAPER = """# 用电量预测模型

## 摘要

本文建立了预测模型，准确率为 92.5%。

**关键词**：时间序列；回归

## 一、问题重述

### 问题背景

背景说明，费用 50%。

## 六、灵敏度分析

### 灵敏度分析

系数 $\\beta_1 = $ 0.8027，截距 $\\beta_0 = $ -0.3396。

![图1 年度用电量](outputs/figures/F1.png)

*图1 年度用电量*

如图1所示。
"""


pytestmark = pytest.mark.usefixtures("offline_kernel", "offline_compiler")

@needs_xelatex
def test_graph_agent_latex_mode_builds_pdf(tmp_path, monkeypatch):
    import ai_system.export.markdown_prep as prep
    def no_pandoc(*args, **kwargs):
        raise AssertionError("native LaTeX must not call pandoc")
    monkeypatch.setattr(prep, "pandoc", no_pandoc)
    ws = _make_workspace(tmp_path)
    events = []

    class Recorder:
        async def send_json_block(self, block_type, content):
            events.append((block_type, content))

    agent = PaperGraphAgent(
        FakeModel(), stream_manager=Recorder(), workspace_dir=str(ws), output_mode="latex"
    )
    output = asyncio.run(agent.run(TASK))

    assert not (ws / "paper.md").exists() and (ws / "paper.tex").exists()
    assert not (workspace_path(ws, '.system/paper_raw.tex')).exists()
    from services.chat_services import paper_project
    from ai_system.research.store import ResearchStore
    store=ResearchStore(ws,agent._load_thread())
    project=paper_project.describe(store)
    assert project["main"]=="paper.tex" and project["compile"]["state"]=="current"
    assert (workspace_path(ws, paper_project.VERSIONS)/project["version"]/"paper.tex").read_bytes()==(ws/"paper.tex").read_bytes()
    assert project["changes"] and all(c["status"]=="applied" for c in project["changes"])
    assert not (workspace_path(ws, '.system/paper_raw.md')).exists()
    assert (ws / "paper.pdf").read_bytes().startswith(b"%PDF")
    assert "paper.pdf" in output
    export_phases = [c["status"] for t, c in events if t == "paper_phase" and c["stage"] == "export"]
    # Events reflect actual export execution, without a synthetic pending phase.
    assert export_phases[0] == "in_progress" and export_phases[-1] == "completed"
    assert "pending" not in export_phases
    assert any(t == "paper_export" and "paper.pdf" in c for t, c in events)
    from services.chat_services.research_session import research_snapshot
    snapshot = research_snapshot(ws, "work", "completed")
    assert snapshot.delivery_stages["export"]["status"] == "completed"
    assert "paper.pdf" in snapshot.artifacts


@needs_pandoc
def test_export_docx_with_default_style(tmp_path):
    from docx import Document
    from PIL import Image
    image = tmp_path / 'outputs/figures/F1.png'
    image.parent.mkdir(parents=True)
    Image.new('RGB', (4, 4), 'white').save(image)

    result = asyncio.run(export_docx(tmp_path, PAPER))
    assert result.reference == "default"
    doc = Document(str(tmp_path / "paper.docx"))
    styles = [p.style.name for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "用电量预测模型" and styles[0] == "Title"
    assert "一、问题重述" in texts and "Heading 1" in styles
    # 公式转成 Word 原生公式；数在公式外的写法先并回公式
    assert "<m:oMath" in doc.element.xml
    # 默认样式：标题不再是 pandoc 的蓝色主题色
    reference = Document(str(tmp_path / '.build/reference.docx'))
    heading = {s.name: s for s in reference.styles}["Heading 1"]
    assert "themeColor" not in heading.element.rPr.xml


MD_CONTRACT = """# 模板契约

## Markdown 骨架
- L1 第1行: 论文题目
- L2 第3行: 摘要
- L2 第5行: 1 研究背景
- L2 第7行: 2 模型假设
- L2 第9行: 3 模型建立与求解
- L3 第10行: 3.1 问题一
- L2 第12行: 4 结论与建议
- L2 第14行: 参考文献
"""


@needs_pandoc
def test_graph_agent_word_mode_follows_template(tmp_path):
    from docx import Document

    ws = _make_workspace(tmp_path)
    (ws / 'template').mkdir(parents=True, exist_ok=True)
    (ws / 'template/contract.md').write_text(MD_CONTRACT, encoding="utf-8")
    # This checks template/export behavior; adversarial prose and its repair are
    # covered by test_graph_end_to_end_offline, including the unregistered number.
    agent = PaperGraphAgent(FakeModel(loose_number_once=False), workspace_dir=str(ws), output_mode="word")
    output = asyncio.run(agent.run(TASK))
    assert "paper.docx" in output

    paper = (ws / "paper.md").read_text(encoding="utf-8")
    headings = [line for line in paper.splitlines() if line.startswith("## ")]
    assert headings == ["## 摘要", "## 1 研究背景", "## 2 模型假设", "## 3 模型建立与求解", "## 4 结论与建议", "## 参考文献"]
    assert "用电量预测" in paper and "政策评估" in paper
    texts = [p.text for p in Document(str(ws / "paper.docx")).paragraphs]
    assert "4 结论与建议" in texts
    from services.chat_services.research_session import research_snapshot
    snapshot = research_snapshot(ws, "work", "completed")
    from ai_system.research.store import ResearchStore
    saved=ResearchStore(ws,snapshot.research_id).snapshot()
    assert [item['title'] for item in saved['publication']['outline']]==[
        '摘要','1 研究背景','2 模型假设','3 模型建立与求解','4 结论与建议','参考文献']
    assert snapshot.delivery_stages["export"]["status"] == "completed"
    assert "paper.docx" in snapshot.artifacts
